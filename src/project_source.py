"""Publish only the factual cells needed by the fixed comparison."""
from pathlib import Path
from decimal import Decimal
import argparse
import csv
import hashlib
import io
import json
import math
import re
import sys
import openpyxl

ROOT = Path(__file__).resolve().parents[1]

def quarter(text):
    match = re.fullmatch(r"(\d{4}):?Q([1-4])", str(text).strip())
    assert match, ("Quarter label", text)
    return int(match[1]) * 4 + int(match[2]) - 1

def label(index):
    return f"{index // 4}Q{index % 4 + 1}"

def vintage(text):
    match = re.fullmatch(r"ROUTPUT(\d{2})Q([1-4])", str(text).strip())
    assert match, ("Vintage header", text)
    year = int(match[1]) + (1900 if int(match[1]) >= 65 else 2000)
    return year * 4 + int(match[2]) - 1

def token(cell):
    value = cell.value
    if value is None:
        return "", "empty"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        assert math.isfinite(value), ("Nonfinite cell", cell.coordinate)
        return str(value), "numeric"
    assert value in ["#N/A", "#VALUE!", "NA", "N/A", "."], ("Unexpected selected token", cell.coordinate, value)
    return str(value), "source_missing"

def csv_bytes(rows, fields):
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue().encode()

def write_json(path, value):
    path.write_bytes((json.dumps(value, indent=2) + "\n").encode())

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--private-dir", type=Path, required=True)
    args = parser.parse_args()
    assert not (ROOT / "data/source_manifest.json").exists(), "Never replace a projection."
    requests = json.loads((args.private_dir / "requests.json").read_text(encoding="utf-8"))
    for request in requests:
        assert request["http_status"] == 200
        assert hashlib.sha256((args.private_dir / request["file"]).read_bytes()).hexdigest() == request["sha256"]
    book = openpyxl.load_workbook(args.private_dir / "quarterly.xlsx", data_only=True)
    sheet = book["ROUTPUT"]
    assert sheet.cell(1, 1).value == "DATE"
    columns = {}
    for col in range(2, sheet.max_column + 1):
        key = vintage(sheet.cell(1, col).value)
        assert key not in columns, "Duplicate vintage header"
        columns[key] = col
    observations = {}
    for row in range(2, sheet.max_row + 1):
        text = sheet.cell(row, 1).value
        if text is None:
            assert all(sheet.cell(row, col).value is None for col in range(1, sheet.max_column + 1))
            continue
        key = quarter(text)
        assert key not in observations, "Duplicate observation quarter"
        observations[key] = row
    origins = list(range(quarter("2000Q1"), quarter("2024Q4") + 1))
    later = quarter("2026Q3")
    assert later in columns and max(columns) == later, "Required latest quarterly vintage"
    needed = set()
    for q in origins:
        for t in range(q - 21, q):
            needed.add((q, t))
            needed.add((later, t))
        for t in [q - 1, q]:
            needed.add((q + 1, t))
            needed.add((later, t))
    levels = []
    for v, t in sorted(needed):
        assert v in columns and t in observations, ("Missing source key", label(v), label(t))
        cell = sheet.cell(observations[t], columns[v])
        value, state = token(cell)
        levels.append(dict(vintage=label(v), observation=label(t), value=value, state=state, sheet=sheet.title, cell=cell.coordinate))
    catalog = []
    for v in sorted({v for v, _ in needed}):
        present = [t for t, r in observations.items() if token(sheet.cell(r, columns[v]))[1] == "numeric"]
        assert present
        catalog.append(dict(vintage=label(v), source_header=sheet.cell(1, columns[v]).value,
                            header_cell=sheet.cell(1, columns[v]).coordinate,
                            last_observation=label(max(present)), source_sheet=sheet.title))
    release_book = openpyxl.load_workbook(args.private_dir / "releases.xlsx", data_only=True)
    notes = " ".join(str(release_book["NOTES"].cell(r, 1).value or "") for r in range(1, 12))
    assert "Q/Q Growth" in notes and "Annual Rate" in notes and "Percentage Points" in notes
    release_sheet = release_book["DATA"]
    headers = []
    for row in range(1, min(15, release_sheet.max_row) + 1):
        for col in range(1, release_sheet.max_column + 1):
            if str(release_sheet.cell(row, col).value).strip().casefold() == "first":
                headers.append((row, col))
    assert len(headers) == 1, ("First column header", headers)
    header_row, first_column = headers[0]
    release_rows = {}
    for row in range(header_row + 1, release_sheet.max_row + 1):
        text = release_sheet.cell(row, 1).value
        if text is None:
            continue
        key = quarter(text)
        assert key not in release_rows, "Duplicate First quarter"
        release_rows[key] = row
    first = []
    for q in origins:
        assert q in release_rows, ("Missing release row", label(q))
        cell = release_sheet.cell(release_rows[q], first_column)
        value, state = token(cell)
        first.append(dict(observation=label(q), value=value, state=state, sheet=release_sheet.title, cell=cell.coordinate))
    data = {
        "data/levels.csv": csv_bytes(levels, ["vintage", "observation", "value", "state", "sheet", "cell"]),
        "data/first_release.csv": csv_bytes(first, ["observation", "value", "state", "sheet", "cell"]),
    }
    for name, content in data.items():
        (ROOT / name).write_bytes(content)
    write_json(ROOT / "data/vintage_catalog.json", catalog)
    manifest = dict(source="Philadelphia Fed RTDSM ROUTPUT, based on BEA releases", requests=requests,
                    planned_origins=len(origins), level_cells=len(levels), first_release_cells=len(first),
                    selected_vintages=len(catalog), revised_vintage=label(later),
                    quarterly_sheet_dimensions=[sheet.max_row, sheet.max_column],
                    releases_header_row=header_row, releases_first_column=first_column,
                    file_sha256={name:hashlib.sha256(content).hexdigest() for name, content in data.items()},
                    numeric_projection="Runtime decimal rendering of stored Excel numeric cells; preserve explicit missing tokens and coordinates",
                    quarterly_information_time="Source-defined middle-quarter snapshot; no exact-day or intraday claim",
                    runtime=dict(python=sys.version.split()[0], openpyxl=openpyxl.__version__),
                    rights="Minimal attributed factual projections only; original workbooks, documents and unused content excluded. See DATA_RIGHTS.md")
    manifest["file_sha256"]["data/vintage_catalog.json"] = hashlib.sha256((ROOT / "data/vintage_catalog.json").read_bytes()).hexdigest()
    write_json(ROOT / "data/source_manifest.json", manifest)
    book.close()
    release_book.close()
    print(json.dumps(dict(status="PROJECTED_NOT_YET_RECONCILED", cells=len(levels)+len(first), vintages=len(catalog), planned_origins=len(origins))))

if __name__ == "__main__":
    main()
